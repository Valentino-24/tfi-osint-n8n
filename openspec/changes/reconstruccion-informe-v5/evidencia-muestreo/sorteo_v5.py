# -*- coding: utf-8 -*-
"""Sorteo reproducible y generacion del material de anotacion V5 (N-08).

Reproduce las cifras del corpus B5 (corte 2026-10-17):
  estrato 1 (positivos del clasificador) = 170
  estrato 2 ('No relevante')             = 554  -> N0 publicado

Genera:
  - muestra_v5_etiquetado.csv/.xlsx   hoja ciega completa (170 + 182 = 352), MEZCLADA,
                                      SIN estrato_origen (no revela la prediccion)
  - muestra_v5_anot1.csv/.xlsx        submuestra comun de 50 (companero 1)
  - muestra_v5_anot2.csv/.xlsx        submuestra comun de 50 (companero 2)
  - muestra_v5_correspondencia.csv    codigo <-> id/estrato/nlp (NO se publica)

Semilla de sorteo: 20261020.  Orden: md5(id || semilla).
Semilla de presentacion: 20261020 (shuffle determinista).
Solo lectura sobre la base (SELECT).
"""
import subprocess, csv, hashlib, io, os, random
from collections import Counter

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

SEMILLA = "20261020"
CORTE  = "2026-10-18 03:00:00+00"
TRACK  = "2026-10-19 00:00:00+00"
N0_MUESTRA = 182
TEXTO_MAX = 3000
OUT = os.path.dirname(os.path.abspath(__file__))

CATEGORIAS = [
    "No relevante", "Phishing", "Robo de Credenciales", "Malware", "Ransomware",
    "Vulnerabilidades", "Filtración de Datos", "Infraestructura y Ataques",
    "Hacktivismo", "Ingeniería Social",
]

SEP = "@@SEP@@"
PSQL = ["docker", "exec", "tfi-postgres", "psql", "-U", "postgres", "-d", "tesi_osint", "-t", "-A", "-F", SEP]

def q(sql):
    r = subprocess.run(PSQL + ["-c", sql], capture_output=True, text=True, encoding="utf-8")
    if r.returncode != 0:
        raise RuntimeError(r.stderr[:600])
    return [line.split(SEP) for line in r.stdout.splitlines() if line.strip()]

Q1 = f"""SELECT p.id, p.title, COALESCE(p.selftext,''), s.display_name
FROM posts p JOIN subreddits s ON s.id = p.subreddit_id
WHERE p.nlp_category <> 'No relevante' AND p.nlp_category IS NOT NULL
  AND p.created_utc < '{CORTE}' AND p.ingested_at < '{TRACK}'
ORDER BY md5(p.id || '{SEMILLA}');"""

Q2 = f"""SELECT p.id, p.title, COALESCE(p.selftext,''), s.display_name
FROM posts p JOIN subreddits s ON s.id = p.subreddit_id
WHERE p.nlp_category = 'No relevante'
  AND p.created_utc < '{CORTE}' AND p.ingested_at < '{TRACK}'
ORDER BY md5(p.id || '{SEMILLA}')
LIMIT {N0_MUESTRA};"""

Q3 = f"""SELECT p.id, p.nlp_category, p.nlp_score
FROM posts p
WHERE p.created_utc < '{CORTE}' AND p.ingested_at < '{TRACK}';"""


def limpiar(t):
    t = (t or "").replace("\r\n", "\n").replace("\r", "\n")
    t = "\n".join(l.rstrip() for l in t.split("\n"))
    while "\n\n\n" in t:
        t = t.replace("\n\n\n", "\n\n")
    return t.strip()[:TEXTO_MAX]


def main():
    e1 = q(Q1)
    e2 = q(Q2)
    meta = {r[0]: (r[1], r[2]) for r in q(Q3)}
    print(f"Estrato 1 (censo positivos): {len(e1)}")
    print(f"Estrato 2 (muestra negativos): {len(e2)} de 554")

    registros = []
    for pid, titulo, texto, com in e1:
        registros.append({"id": pid, "estrato": 1, "titulo": limpiar(titulo),
                          "texto": limpiar(texto), "comunidad": com})
    for pid, titulo, texto, com in e2:
        registros.append({"id": pid, "estrato": 2, "titulo": limpiar(titulo),
                          "texto": limpiar(texto), "comunidad": com})

    # Presentacion mezclada y determinista: el codigo no revela el estrato.
    random.seed(int(SEMILLA))
    random.shuffle(registros)
    for i, r in enumerate(registros, 1):
        r["codigo"] = f"C-{i:03d}"

    # Submuestra comun: 25 del censo + 25 de la muestra.
    r1 = [r for r in registros if r["estrato"] == 1]
    r2 = [r for r in registros if r["estrato"] == 2]
    random.seed(int(SEMILLA) + 1)
    sub = random.sample(r1, 25) + random.sample(r2, 25)
    print(f"Submuestra comun: {len(sub)} (25 + 25)")

    cols = ["codigo_interno", "titulo", "texto", "comunidad", "categoria_asignada", "comentario"]

    def row(r):
        return {"codigo_interno": r["codigo"], "titulo": r["titulo"], "texto": r["texto"],
                "comunidad": r["comunidad"], "categoria_asignada": "", "comentario": ""}

    def escribir_csv(path, filas):
        with io.open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=cols)
            w.writeheader()
            w.writerows(filas)

    def escribir_xlsx(path, filas, titulo_hoja):
        wb = Workbook()
        ws = wb.active
        ws.title = titulo_hoja
        hdr_fill = PatternFill("solid", fgColor="1F3864")
        hdr_font = Font(bold=True, color="FFFFFF", size=11)
        fill_amarillo = PatternFill("solid", fgColor="FFF2CC")
        thin = Side(style="thin", color="D0D0D0")
        borde = Border(left=thin, right=thin, top=thin, bottom=thin)
        anchos = {"codigo_interno": 14, "titulo": 48, "texto": 90,
                  "comunidad": 14, "categoria_asignada": 30, "comentario": 30}
        for j, c in enumerate(cols, 1):
            cel = ws.cell(row=1, column=j, value=c)
            cel.fill = hdr_fill; cel.font = hdr_font
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cel.border = borde
            ws.column_dimensions[get_column_letter(j)].width = anchos[c]
        for i, fila in enumerate(filas, 2):
            for j, c in enumerate(cols, 1):
                cel = ws.cell(row=i, column=j, value=fila[c])
                cel.border = borde
                cel.alignment = Alignment(vertical="top", wrap_text=True)
                if c == "categoria_asignada":
                    cel.fill = fill_amarillo
            ws.row_dimensions[i].height = 60
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(filas)+1}"
        dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIAS) + '"',
                            allow_blank=True, showDropDown=False)
        dv.error = "Elegir una de las 10 categorias de la guia."
        dv.errorTitle = "Categoria no valida"
        ws.add_data_validation(dv)
        dv.add(f"E{2}:E{len(filas)+1}")
        wb.save(path)

    def sha(p):
        return hashlib.sha256(io.open(p, "rb").read()).hexdigest()

    # --- hoja ciega completa ---
    completa = [row(r) for r in registros]
    p_csv = os.path.join(OUT, "muestra_v5_etiquetado.csv")
    p_xlsx = os.path.join(OUT, "muestra_v5_etiquetado.xlsx")
    escribir_csv(p_csv, completa)
    if Workbook:
        escribir_xlsx(p_xlsx, completa, "Etiquetado")

    # --- submuestra comun (una copia por anotador, orden distinto) ---
    for k in (1, 2):
        copia = sub[:]
        random.seed(int(SEMILLA) + 10 + k)
        random.shuffle(copia)
        filas = [row(r) for r in copia]
        pc = os.path.join(OUT, f"muestra_v5_anot{k}.csv")
        px = os.path.join(OUT, f"muestra_v5_anot{k}.xlsx")
        escribir_csv(pc, filas)
        if Workbook:
            escribir_xlsx(px, filas, f"Anotador {k}")

    # --- correspondencia (NO publicar) ---
    corr = os.path.join(OUT, "muestra_v5_correspondencia.csv")
    with io.open(corr, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["codigo_interno", "estrato_origen", "id_plataforma", "nlp_category", "nlp_score"])
        for r in registros:
            cat, score = meta.get(r["id"], ("", ""))
            w.writerow([r["codigo"], r["estrato"], r["id"], cat, score])

    print("\nArchivos e hashes:")
    for p in [p_csv, p_xlsx if Workbook else None, corr,
              os.path.join(OUT, "muestra_v5_anot1.csv"), os.path.join(OUT, "muestra_v5_anot1.xlsx") if Workbook else None,
              os.path.join(OUT, "muestra_v5_anot2.csv"), os.path.join(OUT, "muestra_v5_anot2.xlsx") if Workbook else None]:
        if p and os.path.exists(p):
            print(f"  {os.path.basename(p):34s} sha256={sha(p)}")
    print("\nComunidades hoja completa:", dict(Counter(r["comunidad"] for r in registros)))
    print("Comunidades submuestra   :", dict(Counter(r["comunidad"] for r in sub)))


if __name__ == "__main__":
    main()
