# -*- coding: utf-8 -*-
"""Genera la plantilla de adjudicacion (consenso) y, si estan los 3 archivos
llenos, calcula el kappa de Fleiss y deja solo los desacuerdos.

Uso:
  python evaluar_v5.py                 -> solo plantilla vacia
  python evaluar_v5.py                (tras llenar anot1/anot2 y etiquetado)

Archivo de consenso: adjudicacion_v5.xlsx
"""
import io, csv, os, sys
from collections import Counter, defaultdict

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None

OUT = os.path.dirname(os.path.abspath(__file__))
CATS = [
    "No relevante", "Phishing", "Robo de Credenciales", "Malware", "Ransomware",
    "Vulnerabilidades", "Filtración de Datos", "Infraestructura y Ataques",
    "Hacktivismo", "Ingeniería Social",
]
# aliases
ALIASES = {
    "no relevante": "No relevante",
    "no relev": "No relevante",
    "irrelevante": "No relevante",
    "not relevant": "No relevante",
    "n/a": "No relevante",
}
COLS = ["codigo_interno", "titulo", "etiqueta_anot1", "etiqueta_anot2",
        "etiqueta_anot3", "etiqueta_final", "motivo"]


def kappa_fleiss(matriz):
    """matriz: lista de listas de conteos por categoria (n categorias).
    Devuelve (kappa, P_o, P_e)."""
    n = len(matriz)
    if n == 0:
        return None, None, None
    k = len(matriz[0])
    N = sum(sum(fila) for fila in matriz)          # total de juicios
    if N == 0:
        return None, None, None
    # P_i por sujeto
    P = []
    for fila in matriz:
        s = sum(x * (x - 1) for x in fila)
        P.append(s / (N_cat(fila) * (N_cat(fila) - 1)) if N_cat(fila) > 1 else 0.0)
    P_o = sum(P) / n
    # p_j
    totales = [sum(fila[j] for fila in matriz) for j in range(k)]
    pj = [t / N for t in totales]
    P_e = sum(p * p for p in pj)
    kappa = (P_o - P_e) / (1 - P_e) if (1 - P_e) != 0 else None
    return kappa, P_o, P_e


def N_cat(fila):
    return sum(fila)


def main():
    # 1) Plantilla vacia
    if Workbook:
        wb = Workbook()
        ws = wb.active
        ws.title = "Adjudicacion"
        hdr = PatternFill("solid", fgColor="7F6000")
        hfont = Font(bold=True, color="FFFFFF")
        thin = Side(style="thin", color="D0D0D0")
        borde = Border(left=thin, right=thin, top=thin, bottom=thin)
        anchos = [14, 48, 20, 20, 20, 20, 40]
        for j, c in enumerate(COLS, 1):
            cel = ws.cell(1, j, c); cel.fill = hdr; cel.font = hfont
            cel.alignment = Alignment(horizontal="center", wrap_text=True); cel.border = borde
            ws.column_dimensions[get_column_letter(j)].width = anchos[j-1]
        ws.freeze_panes = "A2"
        ws2 = wb.create_sheet("Instrucciones")
        ws2.column_dimensions["A"].width = 110
        pasos = [
            "ADJUDICACION DE DESACUERDOS - VALIDACION V5",
            "",
            "1. Completar PRIMERO muestra_v5_anot1 / anot2 (companeros) y muestra_v5_etiquetado (autor 3).",
            "2. Enviar los 3 archivos al agente.",
            "3. El agente calcula el kappa de Fleiss sobre los 50 comunes y completa esta planilla",
            "   SOLO con los casos donde los 3 no coinciden.",
            "4. Los 3 autores discuten SOLO esos casos y completan etiqueta_final y motivo.",
            "5. Devolver esta planilla al agente para el calculo de metricas.",
            "",
            "Regla: la etiqueta_final se decide por consenso humano. NUNCA por IA.",
            "Si no hay consenso, decide el criterio mas conservador (No relevante).",
        ]
        for i, t in enumerate(pasos, 1):
            cel = ws2.cell(i, 1, t)
            if i == 1: cel.font = Font(bold=True, size=13)
        wb.save(os.path.join(OUT, "adjudicacion_v5.xlsx"))
        print("Plantilla: adjudicacion_v5.xlsx")

    # 2) Si hay archivos llenos, calcular kappa
    a1 = os.path.join(OUT, "muestra_v5_anot1.csv")
    a2 = os.path.join(OUT, "muestra_v5_anot2.csv")
    full = os.path.join(OUT, "muestra_v5_etiquetado.csv")
    if all(os.path.exists(p) for p in (a1, a2, full)):
        def leer(p):
            res = {}
            for r in csv.DictReader(io.open(p, encoding="utf-8-sig")):
                v = (r["categoria_asignada"] or "").strip()
                if not v:
                    res[r["codigo_interno"]] = v
                    continue
                vl = v.lower()
                if v in CATS:
                    res[r["codigo_interno"]] = v
                elif vl in ALIASES:
                    res[r["codigo_interno"]] = ALIASES[vl]
                elif vl == 'no relevante':
                    res[r["codigo_interno"]] = 'No relevante'
                else:
                    res[r["codigo_interno"]] = v
            return res
        d1, d2, d3 = leer(a1), leer(a2), leer(full)
        comunes = sorted(set(d1) & set(d2) & set(d3))
        llenos = [c for c in comunes if d1[c] and d2[c] and d3[c]]
        print(f"Casos comunes con 3 etiquetas: {len(llenos)} de {len(comunes)}")
        if llenos:
            idx = {c: i for i, c in enumerate(CATS)}
            invalidas = [v for d in (d1, d2, d3) for c in llenos for v in [d.get(c, "")] if v and v not in idx]
            if invalidas:
                print("ATENCION: etiquetas fuera de la lista:", Counter(invalidas))
            matriz = []
            for c in llenos:
                fila = [0] * len(CATS)
                for d in (d1, d2, d3):
                    if d[c] in idx:
                        fila[idx[d[c]]] += 1
                matriz.append(fila)
            k, po, pe = kappa_fleiss(matriz)
            print(f"Kappa de Fleiss = {k:.3f}  (P_o={po:.3f}, P_e={pe:.3f})")
            print("Interpretacion: <0,21 leve | 0,21-0,40 aceptable | 0,41-0,60 moderado | "
                  "0,61-0,80 sustancial | >0,80 casi perfecto")
            # desacuerdos
            des = [c for c in llenos if len({d1[c], d2[c], d3[c]}) > 1]
            print(f"Desacuerdos: {len(des)} casos")
            if Workbook:
                wb = load_workbook(os.path.join(OUT, "adjudicacion_v5.xlsx"))
                ws = wb["Adjudicacion"]
                # limpiar datos previos
                if ws.max_row > 1:
                    ws.delete_rows(2, ws.max_row - 1)
                # titulos desde la hoja completa
                tit = {r["codigo_interno"]: r["titulo"] for r in csv.DictReader(io.open(full, encoding="utf-8-sig"))}
                for i, c in enumerate(des, 2):
                    ws.cell(i, 1, c)
                    ws.cell(i, 2, tit.get(c, ""))
                    ws.cell(i, 3, d1[c]); ws.cell(i, 4, d2[c]); ws.cell(i, 5, d3[c])
                    ws.cell(i, 6, "")
                    ws.cell(i, 7, "")
                wb.save(os.path.join(OUT, "adjudicacion_v5.xlsx"))
                print("adjudicacion_v5.xlsx actualizado con los desacuerdos.")


if __name__ == "__main__":
    main()
