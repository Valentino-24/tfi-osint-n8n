# -*- coding: utf-8 -*-
"""Analisis de la validacion V5 del clasificador por diccionario.

Entradas (en este mismo directorio):
  muestra_v5_etiquetado.csv   -> 352 casos (censo 170 + muestra 182), anotador 3
  muestra_v5_anot1.csv        -> 50 comunes, anotador 1
  muestra_v5_anot2.csv        -> 50 comunes, anotador 2
  adjudicacion_v5.xlsx        -> etiqueta_final/motivo solo para desacuerdos
  muestra_v5_correspondencia.csv -> codigo_interno <-> estrato/nlp (NO se publica)

Salidas:
  resultados_v5.json          -> todas las cifras
  resultados_v5.md            -> resumen legible

Metodo (diseno-muestreo.md):
  - Kappa de Fleiss sobre los 50 comunes (acuerdo inter-anotador, a ciegas).
  - Estrato 1 (N=170) censado: precision = TP/(TP+FP), sin ponderar.
  - Estrato 2 (N=554) muestreado n0=182: sensibilidad = TP/(TP + w*k), w = N0/n0.
  - IC 95 % exactos de Clopper-Pearson.
  - La sensibilidad es un cociente de estimadores: el IC se obtiene transformando
    el IC de la proporcion de omision k/n0 del estrato 2.
"""
import csv, io, os, json
from collections import Counter, defaultdict

OUT = os.path.dirname(os.path.abspath(__file__))

CATS = [
    "No relevante", "Phishing", "Robo de Credenciales", "Malware", "Ransomware",
    "Vulnerabilidades", "Filtración de Datos", "Infraestructura y Ataques",
    "Hacktivismo", "Ingeniería Social",
]
ALIASES = {
    "no relevante": "No relevante",
    "no relev": "No relevante",
    "irrelevante": "No relevante",
    "not relevant": "No relevante",
    "n/a": "No relevante",
}
N1 = 170          # poblacion estrato 1 (censo)
N0 = 554          # poblacion estrato 2 (muestreado)
N0_MUESTRA = 182  # n0


def norm(v):
    v = (v or "").strip()
    if not v:
        return ""
    if v in CATS:
        return v
    vl = v.lower()
    if vl in ALIASES:
        return ALIASES[vl]
    return v


def leer(p):
    return {r["codigo_interno"]: norm(r.get("categoria_asignada"))
            for r in csv.DictReader(io.open(p, encoding="utf-8-sig"))}


# ---------------------------------------------------------------- estadistica
def _betacf(a, b, x):
    """Fraccion continua de Lentz para la beta incompleta."""
    MAXIT, EPS, FPMIN = 300, 3.0e-14, 1.0e-300
    qab, qap, qam = a + b, a + 1.0, a - 1.0
    c = 1.0
    d = 1.0 - qab * x / qap
    if abs(d) < FPMIN:
        d = FPMIN
    d = 1.0 / d
    h = d
    for m in range(1, MAXIT + 1):
        m2 = 2 * m
        aa = m * (b - m) * x / ((qam + m2) * (a + m2))
        d = 1.0 + aa * d
        if abs(d) < FPMIN:
            d = FPMIN
        c = 1.0 + aa / c
        if abs(c) < FPMIN:
            c = FPMIN
        d = 1.0 / d
        h *= d * c
        aa = -(a + m) * (qab + m) * x / ((a + m2) * (qap + m2))
        d = 1.0 + aa * d
        if abs(d) < FPMIN:
            d = FPMIN
        c = 1.0 + aa / c
        if abs(c) < FPMIN:
            c = FPMIN
        d = 1.0 / d
        de = d * c
        h *= de
        if abs(de - 1.0) < EPS:
            break
    return h


def betainc(a, b, x):
    """Beta incompleta regularizada I_x(a,b)."""
    if x <= 0.0:
        return 0.0
    if x >= 1.0:
        return 1.0
    lbeta = (math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b)
             + a * math.log(x) + b * math.log1p(-x))
    front = math.exp(lbeta)
    if x < (a + 1.0) / (a + b + 2.0):
        return front * _betacf(a, b, x) / a
    return 1.0 - front * _betacf(b, a, 1.0 - x) / b


def betainv(p, a, b, lo=0.0, hi=1.0):
    """Inversa de la beta incompleta regularizada por biseccion."""
    for _ in range(200):
        mid = 0.5 * (lo + hi)
        if betainc(a, b, mid) < p:
            lo = mid
        else:
            hi = mid
    return 0.5 * (lo + hi)


def clopper_pearson(x, n, alpha=0.05):
    """IC exacto de una proporcion binomial."""
    if n == 0:
        return (0.0, 1.0)
    lo = 0.0 if x == 0 else betainv(alpha / 2.0, x, n - x + 1)
    hi = 1.0 if x == n else betainv(1.0 - alpha / 2.0, x + 1, n - x)
    return (lo, hi)


import math  # noqa: E402  (despues de las funciones que lo usan en tiempo de ejecucion)


def fleiss(matriz):
    """matriz: filas = sujetos, columnas = categorias, celdas = conteo de jueces."""
    n = len(matriz)
    if n == 0:
        return None, None, None
    k = len(matriz[0])
    N = sum(sum(f) for f in matriz)
    if N == 0:
        return None, None, None
    P = []
    for f in matriz:
        nc = sum(f)
        if nc <= 1:
            P.append(0.0)
            continue
        P.append(sum(x * (x - 1) for x in f) / (nc * (nc - 1)))
    Po = sum(P) / n
    totales = [sum(matriz[i][j] for i in range(n)) for j in range(k)]
    pj = [t / N for t in totales]
    Pe = sum(p * p for p in pj)
    kappa = (Po - Pe) / (1 - Pe) if (1 - Pe) != 0 else None
    return kappa, Po, Pe


def interpretar_kappa(k):
    if k is None:
        return "no calculable"
    if k < 0.21:
        return "leve"
    if k < 0.41:
        return "aceptable"
    if k < 0.61:
        return "moderado"
    if k < 0.81:
        return "sustancial"
    return "casi perfecto"


# ---------------------------------------------------------------- carga
def cargar_adjudicacion():
    adj = {}
    ruta = os.path.join(OUT, "adjudicacion_v5.xlsx")
    if not os.path.exists(ruta):
        return adj
    try:
        import openpyxl
    except ImportError:
        return adj
    wb = openpyxl.load_workbook(ruta)
    if "Adjudicacion" not in wb.sheetnames:
        return adj
    ws = wb["Adjudicacion"]
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if not r or not r[0]:
            continue
        final = norm(r[5]) if len(r) > 5 else ""
        motivo = (r[6] or "") if len(r) > 6 else ""
        if final:
            adj[r[0]] = {"final": final, "motivo": motivo,
                         "anot1": norm(r[2]), "anot2": norm(r[3]), "anot3": norm(r[4])}
    return adj


def main():
    d1 = leer(os.path.join(OUT, "muestra_v5_anot1.csv"))
    d2 = leer(os.path.join(OUT, "muestra_v5_anot2.csv"))
    d3 = leer(os.path.join(OUT, "muestra_v5_etiquetado.csv"))

    with io.open(os.path.join(OUT, "muestra_v5_correspondencia.csv"), encoding="utf-8-sig") as f:
        corr = {r["codigo_interno"]: r for r in csv.DictReader(f)}

    adj = cargar_adjudicacion()
    comunes = sorted(set(d1) & set(d2) & set(d3))
    completos = [c for c in comunes if d1[c] and d2[c] and d3[c]]

    # ---- kappa de Fleiss (acuerdo bruto entre los 3, antes de consenso)
    idx = {c: i for i, c in enumerate(CATS)}
    matriz = []
    for c in completos:
        fila = [0] * len(CATS)
        for d in (d1, d2, d3):
            if d[c] in idx:
                fila[idx[d[c]]] += 1
        matriz.append(fila)
    kappa, Po, Pe = fleiss(matriz)

    desacuerdos = [c for c in completos if len({d1[c], d2[c], d3[c]}) > 1]
    resueltos = [c for c in desacuerdos if c in adj]
    pendientes = [c for c in desacuerdos if c not in adj]

    # ---- estandar de referencia (gold)
    #     1) etiqueta_final del consenso, cuando existe
    #     2) si no, la etiqueta del anotador 3 (autor)
    gold = {}
    for c, v in d3.items():
        if c in adj:
            gold[c] = adj[c]["final"]
        else:
            gold[c] = v

    # ---- matriz de confusion por estrato
    e1 = [c for c in gold if corr.get(c, {}).get("estrato_origen") == "1"]
    e2 = [c for c in gold if corr.get(c, {}).get("estrato_origen") == "2"]

    tp = [c for c in e1 if gold[c] != "No relevante"]
    fp = [c for c in e1 if gold[c] == "No relevante"]
    fn = [c for c in e2 if gold[c] != "No relevante"]
    tn = [c for c in e2 if gold[c] == "No relevante"]

    n0 = len(e2)
    w = N0 / n0 if n0 else None

    precision = len(tp) / len(e1) if e1 else None
    sens = len(tp) / (len(tp) + w * len(fn)) if (len(tp) + w * len(fn)) else None
    f1 = (2 * precision * sens / (precision + sens)) if precision and sens else None

    # IC de la precision (binomial censado)
    p_lo, p_hi = clopper_pearson(len(tp), len(e1))

    # IC de la sensibilidad: transformar el IC de la omision k/n0
    if n0 and len(tp):
        miss = len(fn) / n0
        m_lo, m_hi = clopper_pearson(len(fn), n0)
        # sens = TP/(TP + N0*miss)  (pues w*len(fn) = N0*miss)
        s_hi = len(tp) / (len(tp) + N0 * m_lo) if (len(tp) + N0 * m_lo) else None
        s_lo = len(tp) / (len(tp) + N0 * m_hi) if (len(tp) + N0 * m_hi) else None
    else:
        miss = m_lo = m_hi = s_lo = s_hi = None

    # IC del F1 por propagacion de los extremos (conservador)
    f1_lo = f1_hi = None
    if None not in (p_lo, p_hi, s_lo, s_hi):
        def f1_de(p, s):
            return 2 * p * s / (p + s) if (p + s) else 0.0
        f1_lo = f1_de(p_lo, s_lo)
        f1_hi = f1_de(p_hi, s_hi)

    # ---- por comunidad (precision sobre el censo y omision sobre la muestra)
    com_map = {r["codigo_interno"]: r.get("comunidad", "")
               for r in csv.DictReader(io.open(os.path.join(OUT, "muestra_v5_etiquetado.csv"), encoding="utf-8-sig"))}
    por_com = {}
    for com in sorted(set(com_map.get(c, "") for c in gold)):
        pos = [c for c in e1 if com_map.get(c, "") == com]
        neg = [c for c in e2 if com_map.get(c, "") == com]
        tp_c = sum(1 for c in pos if gold[c] != "No relevante")
        fn_c = sum(1 for c in neg if gold[c] != "No relevante")
        por_com[com] = {
            "censo_n": len(pos),
            "censo_tp": tp_c,
            "censo_fp": len(pos) - tp_c,
            "precision": (tp_c / len(pos)) if pos else None,
            "muestra_n": len(neg),
            "muestra_omitidos": fn_c,
            "omision": (fn_c / len(neg)) if neg else None,
        }

    # ---- acierto exacto de categoria (solo casos clasificados como amenaza por ambos)
    acierto_cat_n = sum(1 for c in gold
                        if gold[c] != "No relevante"
                        and gold[c] == corr.get(c, {}).get("nlp_category"))
    amenazas_ambos = sum(1 for c in gold
                         if gold[c] != "No relevante"
                         and corr.get(c, {}).get("nlp_category") not in ("No relevante", "", None))

    # ---- matriz por categoria (gold x nlp)
    mat_cat = defaultdict(lambda: defaultdict(int))
    for c in gold:
        g = gold[c]
        n = corr.get(c, {}).get("nlp_category", "")
        if g and n:
            mat_cat[g][n] += 1

    # ---- estimacion poblacional de amenazas
    amenazas_est = len(tp) + len(fn) * w if n0 else None

    res = {
        "n_total_etiquetado": len(gold),
        "comunes": len(comunes),
        "comunes_completos": len(completos),
        "desacuerdos": len(desacuerdos),
        "desacuerdos_resueltos": len(resueltos),
        "desacuerdos_pendientes": len(pendientes),
        "kappa_fleiss": kappa,
        "kappa_Po": Po,
        "kappa_Pe": Pe,
        "kappa_interpretacion": interpretar_kappa(kappa),
        "estrato1_n": len(e1),
        "estrato1_tp": len(tp),
        "estrato1_fp": len(fp),
        "estrato2_n": n0,
        "estrato2_fn": len(fn),
        "estrato2_tn": len(tn),
        "peso_w": w,
        "precision": precision,
        "precision_ic": [p_lo, p_hi],
        "sensibilidad": sens,
        "sensibilidad_ic": [s_lo, s_hi],
        "f1": f1,
        "f1_ic": [f1_lo, f1_hi],
        "omision_estrato2": miss,
        "omision_ic": [m_lo, m_hi],
        "acierto_categoria": (acierto_cat_n / amenazas_ambos) if amenazas_ambos else None,
        "acierto_categoria_n": acierto_cat_n,
        "acierto_categoria_base": amenazas_ambos,
        "amenazas_estimadas_poblacion": amenazas_est,
        "fdr": (1 - precision) if precision is not None else None,
        "por_comunidad": por_com,
        "matriz_categoria": {k: dict(v) for k, v in mat_cat.items()},
        "dist_gold": dict(Counter(gold.values())),
    }

    with io.open(os.path.join(OUT, "resultados_v5.json"), "w", encoding="utf-8") as f:
        json.dump(res, f, ensure_ascii=False, indent=2)

    # ---- resumen legible
    L = []
    L.append("# Resultados de la validacion V5\n")
    L.append(f"- Casos etiquetados (anotador 3): **{len(gold)}**")
    L.append(f"- Submuestra comun (3 anotadores): **{len(comunes)}**")
    L.append(f"- Kappa de Fleiss: **{kappa:.3f}** ({interpretar_kappa(kappa)})")
    L.append(f"- Desacuerdos: **{len(desacuerdos)}** | resueltos: {len(resueltos)} | pendientes: {len(pendientes)}\n")
    L.append("## Matriz de confusion\n")
    L.append("| Estrato | n | Amenaza real | No relevante |")
    L.append("|---|---|---|---|")
    L.append(f"| 1 (positivos del clasificador, censo) | {len(e1)} | {len(tp)} | {len(fp)} |")
    L.append(f"| 2 (negativos del clasificador, muestra) | {n0} | {len(fn)} | {len(tn)} |\n")
    L.append("## Estimadores\n")
    L.append("| Metrica | Valor | IC 95 % |")
    L.append("|---|---|---|")
    L.append(f"| Precision | {precision*100:.1f} % | [{p_lo*100:.1f}; {p_hi*100:.1f}] |")
    L.append(f"| Sensibilidad | {sens*100:.1f} % | [{s_lo*100:.1f}; {s_hi*100:.1f}] |")
    L.append(f"| F1 | {f1:.3f} | [{f1_lo:.3f}; {f1_hi:.3f}] |")
    L.append(f"| Omision estrato 2 | {miss*100:.1f} % | [{m_lo*100:.1f}; {m_hi*100:.1f}] |\n")
    L.append(f"Peso de muestreo w = N0/n0 = {N0}/{n0} = {w:.3f}\n")
    L.append("## Por comunidad\n")
    L.append("| Comunidad | Censo n | TP | FP | Precision | Muestra n | Omitidos | Omision |")
    L.append("|---|---:|---:|---:|---:|---:|---:|---:|")
    for com, d in sorted(por_com.items()):
        pc = f"{d['precision']*100:.1f} %" if d["precision"] is not None else "—"
        om = f"{d['omision']*100:.1f} %" if d["omision"] is not None else "—"
        L.append(f"| {com or '(sin dato)'} | {d['censo_n']} | {d['censo_tp']} | {d['censo_fp']} | {pc} | {d['muestra_n']} | {d['muestra_omitidos']} | {om} |")
    L.append("")
    L.append(f"## Otros indicadores\n")
    L.append(f"- Tasa de falsos descubrimientos (1 − precision): **{res['fdr']*100:.1f} %**")
    if res["acierto_categoria"] is not None:
        L.append(f"- Acierto exacto de categoria (sobre {res['acierto_categoria_base']} amenazas detectadas por ambos): **{res['acierto_categoria_n']}/{res['acierto_categoria_base']} = {res['acierto_categoria']*100:.1f} %**")
    if res["amenazas_estimadas_poblacion"]:
        L.append(f"- Amenazas estimadas en el corpus: **{res['amenazas_estimadas_poblacion']:.1f}** de 724 ({100*res['amenazas_estimadas_poblacion']/724:.1f} %)")
    L.append(f"- Distribucion de la referencia: {res['dist_gold']}")
    L.append("")
    L.append("## Matriz por categoria (referencia x clasificador)\n")
    cats_nlp = sorted({n for v in res["matriz_categoria"].values() for n in v})
    L.append("| Referencia \\ Clasificador | " + " | ".join(cats_nlp) + " |")
    L.append("|---|" + "---:|" * len(cats_nlp))
    for g in sorted(res["matriz_categoria"]):
        fila = res["matriz_categoria"][g]
        L.append(f"| {g} | " + " | ".join(str(fila.get(n, 0)) for n in cats_nlp) + " |")


    with io.open(os.path.join(OUT, "resultados_v5.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(L) + "\n")

    print("\n".join(L))


if __name__ == "__main__":
    main()
